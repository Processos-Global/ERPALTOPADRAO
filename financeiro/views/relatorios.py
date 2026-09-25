from datetime import datetime, time
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from financeiro.models import (
    AprovacaoTituloFinanceiro,
    ItemRelatorioPagamento,
    RelatorioPagamento,
    TituloPagar,
)
from financeiro.services.numeracao import gerar_numero
from financeiro.services.permissoes import financeiro_acao_required


def _parse_data(valor):
    try:
        return datetime.strptime(valor or "", "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def _intervalo_local(data_inicio, data_fim):
    tz = timezone.get_current_timezone()
    inicio = timezone.make_aware(datetime.combine(data_inicio, time.min), tz)
    fim = timezone.make_aware(datetime.combine(data_fim, time.max), tz)
    return inicio, fim


def _aprovacoes_disponiveis(data_inicio, data_fim):
    inicio, fim = _intervalo_local(data_inicio, data_fim)
    return (
        AprovacaoTituloFinanceiro.objects.filter(
            decisao=AprovacaoTituloFinanceiro.Decisao.APROVADO,
            criado_em__range=(inicio, fim),
            titulo__status=TituloPagar.Status.APROVADO,
            item_relatorio_pagamento__isnull=True,
        )
        .select_related(
            "titulo",
            "titulo__fornecedor",
            "titulo__obra",
            "titulo__pedido",
            "titulo__plano_financeiro",
            "titulo__plano_financeiro__pai",
            "usuario",
        )
        .order_by("criado_em", "id")
    )


def _dados_item(aprovacao):
    titulo = aprovacao.titulo
    apropriacao = ""
    if titulo.plano_financeiro_id:
        apropriacao = titulo.plano_financeiro.caminho

    pedido_numero = ""
    if titulo.pedido_id:
        pedido_numero = str(getattr(titulo.pedido, "numero", "") or getattr(titulo.pedido, "codigo", "") or titulo.pedido_id)

    centro_custo_codigo = ""
    centro_custo_descricao = ""
    if titulo.plano_financeiro_id:
        centro_custo_codigo = titulo.plano_financeiro.codigo or ""
        centro_custo_descricao = titulo.plano_financeiro.nome or ""

    return {
        "aprovacao": aprovacao,
        "titulo": titulo,
        "data_aprovacao": aprovacao.criado_em,
        "conta_numero": titulo.numero,
        "pedido_numero": pedido_numero,
        "obra_nome": str(titulo.obra) if titulo.obra_id else "",
        "beneficiario_nome": titulo.beneficiario_exibicao,
        "beneficiario_documento": titulo.beneficiario_documento or "",
        "descricao": titulo.descricao,
        "especificacao_pagamento": titulo.especificacao_pagamento or "",
        "apropriacao": apropriacao,
        "parcela": titulo.rotulo_parcela if titulo.rotulo_parcela != "—" else "",
        "vencimento": titulo.vencimento,
        "valor": titulo.saldo_aberto,
        "origem": titulo.get_origem_display(),
        "observacao": titulo.observacao or "",
        "documento_numero": titulo.documento_numero or "",
        "competencia": titulo.competencia,
        "centro_custo_codigo": centro_custo_codigo,
        "centro_custo_descricao": centro_custo_descricao,
        "forma_pagamento": titulo.condicao_pagamento or "",
    }


@financeiro_acao_required("VISUALIZAR")
def relatorios_pagamento(request):
    hoje = timezone.localdate()
    inicio_padrao = hoje.replace(day=1)
    data_inicio = _parse_data(request.GET.get("inicio")) or inicio_padrao
    data_fim = _parse_data(request.GET.get("fim")) or hoje

    if data_fim < data_inicio:
        data_inicio, data_fim = data_fim, data_inicio

    aprovacoes = list(_aprovacoes_disponiveis(data_inicio, data_fim))
    itens_preview = [
        {"aprovacao": a, "titulo": a.titulo, "valor": a.titulo.saldo_aberto}
        for a in aprovacoes
    ]
    total_preview = sum((item["valor"] for item in itens_preview), Decimal("0"))

    historico = (
        RelatorioPagamento.objects.select_related("emitido_por")
        .annotate(total_itens=Sum("itens__valor"))
        .order_by("-emitido_em", "-id")[:100]
    )

    return render(
        request,
        "financeiro/relatorios_pagamento.html",
        {
            "data_inicio": data_inicio,
            "data_fim": data_fim,
            "itens_preview": itens_preview,
            "total_preview": total_preview,
            "historico": historico,
        },
    )


@financeiro_acao_required("PAGAR")
@require_POST
@transaction.atomic
def relatorio_pagamento_emitir(request):
    data_inicio = _parse_data(request.POST.get("inicio"))
    data_fim = _parse_data(request.POST.get("fim"))

    if not data_inicio or not data_fim:
        messages.error(request, "Informe o início e o fim do período.")
        return redirect("financeiro:relatorios_pagamento")
    if data_fim < data_inicio:
        messages.error(request, "A data final não pode ser anterior à data inicial.")
        url = reverse("financeiro:relatorios_pagamento")
        return redirect(f"{url}?inicio={data_inicio:%Y-%m-%d}&fim={data_fim:%Y-%m-%d}")

    # Reconsulta dentro da transação para impedir emissão duplicada da mesma aprovação.
    aprovacoes = list(_aprovacoes_disponiveis(data_inicio, data_fim).select_for_update())
    if not aprovacoes:
        messages.warning(request, "O período pode ser usado novamente, mas não existem novas aprovações ainda não incluídas em relatório. Para rebaixar um relatório anterior, use o histórico de emissões.")
        url = reverse("financeiro:relatorios_pagamento")
        return redirect(f"{url}?inicio={data_inicio:%Y-%m-%d}&fim={data_fim:%Y-%m-%d}")

    relatorio = RelatorioPagamento.objects.create(
        numero=gerar_numero("RELATORIO_PAGAMENTO"),
        periodo_inicio=data_inicio,
        periodo_fim=data_fim,
        emitido_por=request.user,
    )

    itens = []
    total = Decimal("0")
    for aprovacao in aprovacoes:
        dados = _dados_item(aprovacao)
        total += dados["valor"]
        itens.append(ItemRelatorioPagamento(relatorio=relatorio, **dados))

    ItemRelatorioPagamento.objects.bulk_create(itens)
    relatorio.quantidade_itens = len(itens)
    relatorio.valor_total = total
    relatorio.save(update_fields=["quantidade_itens", "valor_total"])

    messages.success(
        request,
        f"Relatório {relatorio.numero} emitido com {len(itens)} pagamento(s).",
    )
    return redirect("financeiro:relatorio_pagamento_excel", pk=relatorio.pk)


@financeiro_acao_required("VISUALIZAR")
def relatorio_pagamento_excel(request, pk):
    relatorio = get_object_or_404(
        RelatorioPagamento.objects.select_related("emitido_por"),
        pk=pk,
    )
    itens = list(relatorio.itens.all().order_by("data_aprovacao", "id"))

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, "A biblioteca openpyxl não está instalada no servidor.")
        return redirect("financeiro:relatorios_pagamento")

    wb = Workbook()
    ws = wb.active
    ws.title = "Programação de pagamentos"
    ws.sheet_view.showGridLines = False

    # Modelo utilizado pelo setor financeiro: uma linha por compromisso/parcela.
    cabecalhos = [
        "ID_Compromisso",
        "Parcela/Evento",
        "Nº da NF",
        "Data_Prevista",
        "Competência",
        "Fornecedor",
        "Centro_Custo",
        "Descrição_CC",
        "Valor_Previsto",
        "Valor_Reprogramado",
        "Valor_Final_Previsto",
        "Status",
        "Forma_Pagamento",
        "Observações",
        "DADOS BENEFICIÁRIO",
    ]

    azul = "1F4E78"
    amarelo = "FFF2CC"
    branco = "FFFFFF"
    azul_texto = "0000FF"
    borda_fina = Side(style="thin", color="000000")

    for coluna, texto in enumerate(cabecalhos, 1):
        celula = ws.cell(row=1, column=coluna, value=texto)
        celula.fill = PatternFill("solid", fgColor=azul)
        celula.font = Font(bold=True, color=branco)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)

    for linha, item in enumerate(itens, 2):
        competencia = item.competencia or item.vencimento
        competencia_texto = competencia.strftime("%b/%y").lower() if competencia else ""
        parcela = item.parcela or "ÚNICA"

        valores = [
            item.conta_numero,
            parcela,
            item.documento_numero,
            item.vencimento,
            competencia_texto,
            item.beneficiario_nome,
            item.centro_custo_codigo,
            item.centro_custo_descricao,
            item.valor,
            None,
            item.valor,
            "Previsto",
            item.forma_pagamento,
            item.observacao,
            item.beneficiario_documento,
        ]

        for coluna, valor in enumerate(valores, 1):
            celula = ws.cell(row=linha, column=coluna, value=valor)
            celula.alignment = Alignment(vertical="center", wrap_text=True)
            celula.border = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)

        # Colunas de programação seguem o visual da planilha de referência.
        for coluna in (1, 2, 3, 4, 9, 10, 12, 13, 14, 15):
            ws.cell(row=linha, column=coluna).fill = PatternFill("solid", fgColor=amarelo)

        for coluna in (1, 2, 3, 4, 9, 12, 13):
            ws.cell(row=linha, column=coluna).font = Font(color=azul_texto)

        ws.cell(row=linha, column=4).number_format = "dd/mm/yyyy"
        ws.cell(row=linha, column=9).number_format = 'R$ #,##0.00'
        ws.cell(row=linha, column=10).number_format = 'R$ #,##0.00'
        ws.cell(row=linha, column=11).number_format = 'R$ #,##0.00'

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{max(1, len(itens) + 1)}"
    ws.row_dimensions[1].height = 28

    larguras = [18, 16, 16, 16, 14, 32, 15, 42, 18, 21, 21, 16, 20, 34, 24]
    for idx, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(idx)].width = largura

    # Segunda aba: rastreabilidade da emissão sem alterar o layout operacional.
    meta = wb.create_sheet("Controle da emissão")
    meta.sheet_view.showGridLines = False
    controle = [
        ("Relatório", relatorio.numero),
        ("Período das aprovações", f"{relatorio.periodo_inicio:%d/%m/%Y} a {relatorio.periodo_fim:%d/%m/%Y}"),
        ("Emitido em", timezone.localtime(relatorio.emitido_em).strftime("%d/%m/%Y %H:%M")),
        ("Emitido por", (relatorio.emitido_por.get_full_name() or relatorio.emitido_por.get_username()) if relatorio.emitido_por else ""),
        ("Quantidade", relatorio.quantidade_itens),
        ("Valor total", relatorio.valor_total),
    ]
    for linha, (rotulo, valor) in enumerate(controle, 1):
        meta.cell(row=linha, column=1, value=rotulo).font = Font(bold=True)
        meta.cell(row=linha, column=2, value=valor)
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 46
    meta.cell(row=6, column=2).number_format = 'R$ #,##0.00'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{relatorio.numero}.xlsx"'
    return response

