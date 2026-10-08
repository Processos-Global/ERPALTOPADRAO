from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.db import transaction
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from usuarios.decorators import modulo_required
from usuarios.models import ModuloSistema, NivelPermissao

from obras.models import MovimentoEstoque, Obra
from obras.services import obter_estoque, obter_item_estoque, resumo_por_obra


def _decimal_positivo(valor):
    texto = str(valor or "").strip().replace(" ", "")
    if not texto:
        return None
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        numero = Decimal(texto)
    except (InvalidOperation, ValueError):
        return None
    return numero if numero > 0 else None


def _nome_obra(obra):
    return (obra.nome_curto or obra.nome or obra.codigo or "Obra").strip()


@modulo_required(ModuloSistema.OBRAS, NivelPermissao.LEITURA)
def almoxarifado_dashboard(request):
    q = (request.GET.get("q") or "").strip().lower()
    obra_id = request.GET.get("obra")

    todas_linhas = obter_estoque(somente_com_saldo=True)
    resumo = resumo_por_obra(todas_linhas)
    valor_total_geral = sum((linha["valor_estoque"] for linha in todas_linhas), Decimal("0"))

    linhas = todas_linhas
    if obra_id and str(obra_id).isdigit():
        linhas = [linha for linha in linhas if linha["obra_id"] == int(obra_id)]
    if q:
        linhas = [
            linha for linha in linhas
            if q in linha["material_nome"].lower()
            or q in linha["descricao"].lower()
            or q in _nome_obra(linha["obra"]).lower()
            or q in (linha["obra"].codigo or "").lower()
        ]

    valor_total_filtrado = sum((linha["valor_estoque"] for linha in linhas), Decimal("0"))

    return render(request, "obras/almoxarifado_dashboard.html", {
        "linhas": linhas,
        "resumo_obras": resumo,
        "obras": Obra.objects.filter(ativa=True).order_by("nome"),
        "obra_filtro": int(obra_id) if obra_id and str(obra_id).isdigit() else None,
        "q": request.GET.get("q", ""),
        "total_itens": len(linhas),
        "valor_total_geral": valor_total_geral,
        "valor_total_filtrado": valor_total_filtrado,
    })


@modulo_required(ModuloSistema.OBRAS, NivelPermissao.LEITURA)
def almoxarifado_obra(request, obra_id):
    obra = get_object_or_404(Obra, pk=obra_id)
    q = (request.GET.get("q") or "").strip().lower()
    estoque_completo = obter_estoque(obra=obra, somente_com_saldo=True)
    valor_total_obra = sum((linha["valor_estoque"] for linha in estoque_completo), Decimal("0"))

    linhas = estoque_completo
    if q:
        linhas = [
            linha for linha in linhas
            if q in linha["material_nome"].lower() or q in linha["descricao"].lower()
        ]

    movimentos = (
        MovimentoEstoque.objects
        .filter(Q(obra_origem=obra) | Q(obra_destino=obra))
        .select_related("obra_origem", "obra_destino", "material", "criado_por")
        .order_by("-data_movimento", "-id")[:100]
    )

    return render(request, "obras/almoxarifado_obra.html", {
        "obra": obra,
        "linhas": linhas,
        "estoque_modal": estoque_completo,
        "destinos": Obra.objects.filter(ativa=True).exclude(pk=obra.pk).order_by("nome"),
        "movimentos": movimentos,
        "q": request.GET.get("q", ""),
        "total_itens": len(estoque_completo),
        "valor_total_obra": valor_total_obra,
    })


@require_POST
@modulo_required(ModuloSistema.OBRAS, NivelPermissao.EDICAO)
def saida_estoque(request, obra_id):
    obra = get_object_or_404(Obra, pk=obra_id, ativa=True)
    chave_selecionada = request.POST.get("item") or ""
    item = obter_item_estoque(obra, chave_selecionada) if chave_selecionada else None

    if item is None:
        messages.error(request, "Selecione um item válido do estoque.")
        return redirect("obras:almoxarifado_obra", obra_id=obra.pk)

    quantidade = _decimal_positivo(request.POST.get("quantidade"))
    if quantidade is None:
        messages.error(request, "Informe uma quantidade válida maior que zero.")
        return redirect("obras:almoxarifado_obra", obra_id=obra.pk)
    if quantidade > item["saldo"]:
        messages.error(
            request,
            f"Saldo insuficiente. Disponível: {item['saldo']:.2f} {item['unidade']}.",
        )
        return redirect("obras:almoxarifado_obra", obra_id=obra.pk)

    with transaction.atomic():
        Obra.objects.select_for_update().get(pk=obra.pk)
        item_atual = obter_item_estoque(obra, chave_selecionada)
        if item_atual is None or quantidade > item_atual["saldo"]:
            messages.error(request, "O saldo foi alterado. Atualize a página e tente novamente.")
            return redirect("obras:almoxarifado_obra", obra_id=obra.pk)

        MovimentoEstoque.objects.create(
            tipo=MovimentoEstoque.Tipo.SAIDA,
            obra_origem=obra,
            material_id=item_atual["material_id"],
            descricao_item=item_atual["descricao"],
            unidade=item_atual["unidade"],
            quantidade=quantidade,
            finalidade=(request.POST.get("finalidade") or "").strip(),
            documento_referencia=(request.POST.get("documento_referencia") or "").strip(),
            observacao=(request.POST.get("observacao") or "").strip(),
            criado_por=request.user,
        )

    messages.success(request, "Saída registrada com sucesso.")
    return redirect("obras:almoxarifado_obra", obra_id=obra.pk)


@require_POST
@modulo_required(ModuloSistema.OBRAS, NivelPermissao.EDICAO)
def transferencia_estoque(request, obra_id):
    obra = get_object_or_404(Obra, pk=obra_id, ativa=True)
    chave_selecionada = request.POST.get("item") or ""
    item = obter_item_estoque(obra, chave_selecionada) if chave_selecionada else None
    destino = Obra.objects.filter(
        ativa=True,
        pk=request.POST.get("obra_destino"),
    ).exclude(pk=obra.pk).first()
    quantidade = _decimal_positivo(request.POST.get("quantidade"))

    if item is None:
        messages.error(request, "Selecione um item válido do estoque.")
    elif destino is None:
        messages.error(request, "Selecione uma obra de destino válida.")
    elif quantidade is None:
        messages.error(request, "Informe uma quantidade válida maior que zero.")
    elif quantidade > item["saldo"]:
        messages.error(
            request,
            f"Saldo insuficiente. Disponível: {item['saldo']:.2f} {item['unidade']}.",
        )
    else:
        with transaction.atomic():
            # Bloqueia as duas obras em ordem fixa para evitar disputa entre transferências simultâneas.
            list(
                Obra.objects.select_for_update()
                .filter(pk__in=sorted([obra.pk, destino.pk]))
                .order_by("pk")
            )
            item_atual = obter_item_estoque(obra, chave_selecionada)
            if item_atual is None or quantidade > item_atual["saldo"]:
                messages.error(request, "O saldo foi alterado. Atualize a página e tente novamente.")
                return redirect("obras:almoxarifado_obra", obra_id=obra.pk)

            MovimentoEstoque.objects.create(
                tipo=MovimentoEstoque.Tipo.TRANSFERENCIA,
                custo_unitario_transferencia=item_atual["valor_unitario_medio"],
                obra_origem=obra,
                obra_destino=destino,
                material_id=item_atual["material_id"],
                descricao_item=item_atual["descricao"],
                unidade=item_atual["unidade"],
                quantidade=quantidade,
                finalidade=(request.POST.get("finalidade") or "").strip(),
                documento_referencia=(request.POST.get("documento_referencia") or "").strip(),
                observacao=(request.POST.get("observacao") or "").strip(),
                criado_por=request.user,
            )
        messages.success(request, f"Transferência para {_nome_obra(destino)} registrada com sucesso.")

    return redirect("obras:almoxarifado_obra", obra_id=obra.pk)


@modulo_required(ModuloSistema.OBRAS, NivelPermissao.LEITURA)
def relatorio_transferencias(request, obra_id):
    """Depreciação individual por movimento, usada somente no Excel de ressarcimento."""
    from datetime import datetime
    from io import BytesIO
    from django.http import HttpResponse
    from django.utils import timezone
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    obra = get_object_or_404(Obra, pk=obra_id)
    movimentos = (MovimentoEstoque.objects
        .filter(tipo=MovimentoEstoque.Tipo.TRANSFERENCIA, obra_origem=obra)
        .select_related("obra_destino", "obra_origem")
        .order_by("data_movimento", "id"))
    inicio = request.GET.get("inicio", "")
    fim = request.GET.get("fim", "")
    destino = request.GET.get("destino", "")
    if inicio:
        try:
            movimentos = movimentos.filter(data_movimento__date__gte=datetime.strptime(inicio, "%Y-%m-%d").date())
        except ValueError:
            return HttpResponse("Data inicial inválida.", status=400)
    if fim:
        try:
            movimentos = movimentos.filter(data_movimento__date__lte=datetime.strptime(fim, "%Y-%m-%d").date())
        except ValueError:
            return HttpResponse("Data final inválida.", status=400)
    if destino:
        if not destino.isdigit():
            return HttpResponse("Destino inválido.", status=400)
        movimentos = movimentos.filter(obra_destino_id=int(destino))

    registros = list(movimentos)
    if request.method != "POST":
        for mov in registros:
            mov.valor_bruto_relatorio = (mov.custo_unitario_transferencia * mov.quantidade
                if mov.custo_unitario_transferencia is not None else None)
        return render(request, "obras/relatorio_transferencias.html", {
            "obra": obra,
            "destinos": Obra.objects.filter(ativa=True).exclude(pk=obra_id).order_by("nome"),
            "inicio": inicio, "fim": fim, "destino": destino,
            "qtd": len(registros), "movimentos": registros,
        })

    if not registros:
        return HttpResponse("Nenhuma transferência selecionada para gerar o relatório.", status=400)
    porcentagens = {}
    for mov in registros:
        campo = f"depreciacao_{mov.pk}"
        texto = request.POST.get(campo)
        if texto is None or not texto.strip():
            return HttpResponse(f"Informe a depreciação do movimento {mov.pk}.", status=400)
        texto = texto.strip().replace(",", ".")
        try:
            valor = Decimal(texto)
        except (ValueError, InvalidOperation):
            return HttpResponse(f"Depreciação inválida no movimento {mov.pk}.", status=400)
        if not valor.is_finite() or not 0 <= valor <= 100:
            return HttpResponse(f"Depreciação do movimento {mov.pk} deve estar entre 0 e 100%.", status=400)
        porcentagens[mov.pk] = valor

    wb = Workbook()
    ws = wb.active
    ws.title = "Transferências"
    ws.append(["RELATÓRIO DE TRANSFERÊNCIAS ENTRE OBRAS"])
    ws.append(["Obra de origem", _nome_obra(obra), "Emitido em", timezone.localtime().strftime("%d/%m/%Y %H:%M")])
    ws.append(["Data", "Movimento", "Obra origem", "Obra destino", "Material", "Un.", "Quantidade",
        "Custo unit. compra (R$)", "Depreciação individual (%)", "Valor unit. depreciado (R$)",
        "Valor a ressarcir (R$)", "Referência", "Observações"])
    total = Decimal("0")
    sem_custo = 0
    for mov in registros:
        depreciacao = porcentagens[mov.pk]
        custo = mov.custo_unitario_transferencia
        if custo is None:
            sem_custo += 1
        unitario = custo * (Decimal("1") - depreciacao / Decimal("100")) if custo is not None else None
        subtotal = unitario * mov.quantidade if unitario is not None else None
        if subtotal is not None:
            total += subtotal
        ws.append([timezone.localtime(mov.data_movimento).strftime("%d/%m/%Y"), mov.pk,
            _nome_obra(mov.obra_origem), _nome_obra(mov.obra_destino), mov.descricao_item,
            mov.unidade, float(mov.quantidade), float(custo) if custo is not None else None,
            float(depreciacao), float(unitario) if unitario is not None else None,
            float(subtotal) if subtotal is not None else None,
            mov.documento_referencia, mov.observacao])
    linha_total = len(registros) + 4
    ws.append(["TOTAL A RESSARCIR", None, None, None, None, None, None, None, None, None, float(total)])
    ws.append(["Movimentos sem custo histórico", sem_custo,
        "Validar o custo de compra antes do pagamento. Movimentos sem custo não estão incluídos no total."])
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:M{len(registros)+3}"
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="453330")
        cell.alignment = Alignment(wrap_text=True)
    for col, largura in {"A":23,"B":14,"C":24,"D":24,"E":45,"F":9,"G":15,"H":25,
                         "I":29,"J":30,"K":27,"L":22,"M":40}.items():
        ws.column_dimensions[col].width = largura
    for row in ws.iter_rows(min_row=4, max_row=3+len(registros)):
        for idx in (8,10,11):
            row[idx-1].number_format = '"R$" #,##0.00'
        row[8].number_format = '0.00"%"'
    ws[f"K{linha_total}"].number_format = '"R$" #,##0.00'
    saida = BytesIO()
    wb.save(saida)
    response = HttpResponse(saida.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="transferencias_obra_{obra.pk}_{timezone.localdate():%Y%m%d}.xlsx"'
    return response
